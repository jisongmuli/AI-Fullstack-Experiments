"""Portable reports for document analysis, without claiming authorship ground truth.

Report annotations apply to extracted text and Python Unicode-codepoint offsets.
They do not modify the source file or reproduce its page layout.
"""
from __future__ import annotations

import io
import json
import math
import re
from html import escape
from typing import Any

REPORT_NOTE = (
    "AI 特征覆盖率表示模型评分达到高风险阈值的非空白字符占比，"
    "不是实际 AI 撰写比例，也不是作者身份或内容来源的证明。"
)
ANNOTATION_NOTE = (
    "下划线按模型评分片段标注提取后的原文；它不表示逐字定位准确率，"
    "也不修改原文件或保留其原始分页、图片、公式和复杂排版。"
)
MIME_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "html": "text/html; charset=utf-8",
    "json": "application/json; charset=utf-8",
}


def _filename(source: Any, report_format: str) -> str:
    name = str(source or "粘贴文本").replace("\\", "/").rsplit("/", 1)[-1]
    name = name.rsplit(".", 1)[0] if "." in name else name
    name = re.sub(r"[^\w .-]", "_", name, flags=re.UNICODE).strip(" ._")[:72]
    return f"文档检测报告_{name or '文本'}.{report_format}"


def _score(value: Any) -> float | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    number = float(value)
    return number if math.isfinite(number) and 0 <= number <= 1 else None


def _percent(value: Any) -> str:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return "未提供"
    number = float(value)
    return f"{number:.1f}%" if math.isfinite(number) and 0 <= number <= 100 else "未提供"


def _score_label(value: Any) -> str:
    number = _score(value)
    return f"{100 * number:.1f}/100" if number is not None else "未评分"


def _validated_segments(result: dict[str, Any]) -> list[dict[str, Any]]:
    text = result.get("extracted_text")
    if not isinstance(text, str):
        raise ValueError("Report requires extracted_text as a string.")
    detection = result.get("detection") or {}
    if not isinstance(detection, dict):
        raise ValueError("Report detection must be an object.")
    segments = detection.get("segments") or []
    if not isinstance(segments, list):
        raise ValueError("Report segments must be a list.")
    validated = []
    previous_end = 0
    for segment in segments:
        if not isinstance(segment, dict):
            raise ValueError("Report segment must be an object.")
        start, end = segment.get("start"), segment.get("end")
        if (type(start) is not int or type(end) is not int
                or not 0 <= start < end <= len(text) or start < previous_end):
            raise ValueError("Report segment offsets are invalid or overlap.")
        if segment.get("text") != text[start:end]:
            raise ValueError("Report segment text does not match extracted text.")
        if segment.get("highlighted") and _score(segment.get("risk_score")) is None:
            raise ValueError("Highlighted report segments require a finite model score.")
        previous_end = end
        validated.append(segment)
    return validated


def _summary_rows(result: dict[str, Any]) -> list[tuple[str, str]]:
    summary = result.get("document_summary") or {}
    detection = result.get("detection") or {}
    return [
        ("AI 特征覆盖率 以全文为分母", _percent(summary.get("highlighted_total_percent"))),
        ("高风险字符占比 以可评分内容为分母", _percent(summary.get("highlighted_character_percent"))),
        ("模型评分覆盖率 以全文为分母", _percent(summary.get("scored_character_percent"))),
        ("整篇模型风险分数", _score_label(detection.get("risk_score"))
         if _score(detection.get("risk_score")) is not None else "未提供整篇评分"),
        ("可评分片段模型均分", _score_label(summary.get("mean_segment_score"))),
        ("高风险片段数", str(summary.get("high_risk_segments", "未提供"))),
        ("片段总数", str(summary.get("total_segments", "未提供"))),
        ("未评分片段数", str(summary.get("unscored_segments", "未提供"))),
    ]


def _highlight_runs(text: str, segments: list[dict[str, Any]], start: int, end: int):
    """Yield exact original spans, optionally associated with a flagged segment."""
    cursor = start
    for segment in segments:
        if not segment.get("highlighted") or segment["end"] <= start:
            continue
        if segment["start"] >= end:
            break
        left, right = max(start, segment["start"]), min(end, segment["end"])
        if cursor < left:
            yield text[cursor:left], None
        if left < right:
            yield text[left:right], segment
        cursor = right
    if cursor < end:
        yield text[cursor:end], None


def _metadata(result: dict[str, Any]) -> list[tuple[str, str]]:
    detection = result.get("detection") or {}
    return [
        ("来源", str(result.get("filename") or "粘贴文本")),
        ("文档格式", str(result.get("format") or "text")),
        ("检测时间", str(result.get("created_at") or "未提供")),
        ("模型版本", str(detection.get("model_version") or "未提供")),
        ("来源 SHA256", str(result.get("source_sha256") or "未提供")),
    ]


def _warnings(result: dict[str, Any]) -> list[str]:
    warnings = result.get("warnings") or []
    return [str(item) for item in warnings if item] if isinstance(warnings, list) else []


def _source_location(result: dict[str, Any], segment: dict[str, Any]) -> str:
    """Human-readable trace to extracted pages/paragraphs, never to visual geometry."""
    sections = result.get("sections") or []
    if not isinstance(sections, list):
        return "提取原文"
    locations = []
    for section in sections:
        if not isinstance(section, dict):
            continue
        left, right = section.get("start"), section.get("end")
        if type(left) is not int or type(right) is not int:
            continue
        if left >= segment["end"] or right <= segment["start"]:
            continue
        label = str(section.get("label") or "原文位置")
        page = section.get("page")
        if type(page) is int and page > 0 and str(page) not in label:
            label = f"第 {page} 页 {label}"
        if label not in locations:
            locations.append(label)
    if not locations:
        return "提取原文"
    short = "、".join(locations[:4])
    return short if len(locations) <= 4 else f"{short} 等 {len(locations)} 个位置"


def _html_report(result: dict[str, Any], segments: list[dict[str, Any]]) -> bytes:
    text = result["extracted_text"]
    summary_rows = "".join(f"<tr><th scope='row'>{escape(k)}</th><td>{escape(v)}</td></tr>"
                           for k, v in _summary_rows(result))
    metadata_rows = "".join(f"<dt>{escape(k)}</dt><dd>{escape(v)}</dd>" for k, v in _metadata(result))
    original = []
    for span, segment in _highlight_runs(text, segments, 0, len(text)):
        if segment is None:
            original.append(escape(span))
        else:
            tooltip = "片段模型分数 " + _score_label(segment.get("risk_score"))
            original.append(f"<span class='flag' title='{escape(tooltip, quote=True)}'>{escape(span)}</span>")
    warnings = "".join(f"<li>{escape(item)}</li>" for item in _warnings(result))
    warning_html = f"<h2>提取和评分提示</h2><ul>{warnings}</ul>" if warnings else ""
    details = []
    for index, segment in enumerate(segments, 1):
        score = _score(segment.get("risk_score"))
        status = "高风险片段" if segment.get("highlighted") else ("已评分" if score is not None else "未评分")
        evidence = segment.get("evidence") or []
        evidence = "; ".join(str(item) for item in evidence) if isinstance(evidence, list) else str(evidence)
        details.append(f"<tr><td>{index}</td><td>{segment['start']}–{segment['end']}</td>"
                       f"<td>{escape(status)}</td><td>{escape(_score_label(score) if score is not None else '未评分')}</td>"
                       f"<td>{escape(_source_location(result, segment))}<br>{escape(evidence)}</td></tr>")
    detail_html = ("<h2>片段评分依据</h2><table class='details'><thead><tr><th>片段</th><th>字符位置</th>"
                   "<th>状态</th><th>模型分数</th><th>来源位置和依据</th></tr></thead><tbody>"
                   + "".join(details) + "</tbody></table>") if details else ""
    html = f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<meta http-equiv="Content-Security-Policy" content="default-src 'none'; style-src 'unsafe-inline'; base-uri 'none'; form-action 'none'">
<title>文档 AI 特征检测报告</title><style>
body{{margin:0;background:#f7f8f8;color:#202c2d;font:16px/1.65 'Microsoft YaHei','Arial',sans-serif}}
main{{max-width:960px;margin:auto;padding:40px;background:#fff}}h1{{font-size:28px;color:#000;margin:0 0 18px}}
h2{{font-size:20px;color:#000;margin:30px 0 12px}}p{{margin:10px 0}}table{{border-collapse:collapse;width:100%;margin:14px 0}}
th,td{{border:1px solid #d9d9d9;padding:10px 12px;text-align:left;vertical-align:middle;overflow-wrap:anywhere}}
thead th{{background:#eef0f2}}table:not(.details) th{{width:65%;font-weight:500}}dl{{display:grid;grid-template-columns:120px 1fr;gap:8px 16px}}
dt{{font-weight:600}}dd{{margin:0;overflow-wrap:anywhere}}.original{{white-space:pre-wrap;overflow-wrap:anywhere;line-height:1.85}}
.flag{{text-decoration:underline;text-decoration-color:#ad5100;text-decoration-thickness:2px;text-underline-offset:4px;background:#fff6db}}
.details{{font-size:14px}}.details td:first-child{{width:40px}}@media(max-width:600px){{main{{padding:20px}}dl{{grid-template-columns:1fr;gap:4px}}dd{{margin-bottom:12px}}}}
@media print{{body{{background:#fff}}main{{max-width:none;padding:0}}h2{{break-after:avoid}}tr{{break-inside:avoid}}thead{{display:table-header-group}}.flag{{background:transparent}}}}
@page{{size:A4;margin:20mm}}
</style></head><body><main><h1>文档 AI 特征检测报告</h1>
<p>{escape(REPORT_NOTE)}</p><p>{escape(ANNOTATION_NOTE)}</p>
<h2>检测结果</h2><table><tbody>{summary_rows}</tbody></table>
<p>均分仅描述已评分片段的模型输出，不是校准后的生成概率。未评分内容没有被记为人类创作。</p>
<h2>来源和可复核信息</h2><dl>{metadata_rows}</dl>{warning_html}
<h2>带标注的提取原文</h2><p>下划线表示达到高风险阈值的模型评分片段。未加下划线不证明内容由人类撰写。</p>
<div class="original">{''.join(original)}</div>{detail_html}
</main></body></html>"""
    return html.encode("utf-8")


def _docx_report(result: dict[str, Any], segments: list[dict[str, Any]]) -> bytes:
    # These imports are application dependencies, not an artifact-authoring runtime.
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.1)
    for style_name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    doc.styles["Title"].font.size = Pt(23)
    doc.styles["Title"].paragraph_format.space_after = Pt(16)
    doc.styles["Heading 1"].font.size = Pt(15)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(16)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(8)
    for style_name in ("Title", "Subtitle", "Heading 1", "Heading 2"):
        paragraph_properties = doc.styles[style_name]._element.get_or_add_pPr()
        for border in list(paragraph_properties.findall(qn("w:pBdr"))):
            paragraph_properties.remove(border)
    doc.core_properties.title = "文档 AI 特征检测报告"
    doc.core_properties.subject = "提取原文的模型评分和高风险片段标注"
    doc.core_properties.author = "内容可信工作台"
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("第 ").font.size = Pt(9)
    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    footer._p.append(page_field)
    footer.add_run(" 页").font.size = Pt(9)
    doc.add_paragraph("文档 AI 特征检测报告", "Title")
    doc.add_paragraph(REPORT_NOTE)
    doc.add_paragraph(ANNOTATION_NOTE)
    doc.add_heading("检测结果", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width, table.columns[1].width = Cm(12.5), Cm(4.3)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "指标", "结果"
    for label, value in _summary_rows(result):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, value
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge_name in ("top", "left", "bottom", "right"):
                edge = OxmlElement(f"w:{edge_name}")
                for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    edge.set(qn(f"w:{key}"), value)
                borders.append(edge)
            properties.append(borders)
            margins = OxmlElement("w:tcMar")
            for side in ("top", "bottom", "left", "right"):
                edge = OxmlElement(f"w:{side}")
                edge.set(qn("w:w"), "100")
                edge.set(qn("w:type"), "dxa")
                margins.append(edge)
            properties.append(margins)
            if row_index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "EEF0F2")
                properties.append(shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                if row_index == 0:
                    for run in paragraph.runs:
                        run.bold = True
    repeat_header = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat_header)
    doc.add_paragraph("均分仅描述已评分片段的模型输出，不是校准后的生成概率。未评分内容没有被记为人类创作。")
    doc.add_heading("来源和可复核信息", level=1)
    for label, value in _metadata(result):
        paragraph = doc.add_paragraph()
        paragraph.add_run(f"{label}：").bold = True
        paragraph.add_run(value)
    warnings = _warnings(result)
    if warnings:
        doc.add_heading("提取和评分提示", level=1)
        for warning in warnings:
            doc.add_paragraph(warning)
    doc.add_heading("带标注的提取原文", level=1)
    doc.add_paragraph("下划线表示达到高风险阈值的模型评分片段。未加下划线不证明内容由人类撰写。")
    text = result["extracted_text"]
    cursor = 0
    # Splitting only on LF keeps blank paragraphs and preserves Python offsets.
    for line in text.split("\n"):
        paragraph = doc.add_paragraph()
        for span, segment in _highlight_runs(text, segments, cursor, cursor + len(line)):
            run = paragraph.add_run(span)
            if segment is not None:
                run.underline = True
                run.font.color.rgb = RGBColor.from_string("9A4300")
        cursor += len(line) + 1
    if segments:
        doc.add_heading("片段评分依据", level=1)
        doc.add_paragraph("字符位置按 Unicode 字符编号，起点包含，终点不包含。下列依据描述统计关联，不是内容来源证明。")
        for index, segment in enumerate(segments, 1):
            score = _score(segment.get("risk_score"))
            status = "高风险片段" if segment.get("highlighted") else ("已评分" if score is not None else "未评分")
            score_label = _score_label(score) if score is not None else "未评分"
            title = doc.add_paragraph(f"片段 {index}  字符位置 {segment['start']} 至 {segment['end']}  {status}  模型分数 {score_label}")
            title.paragraph_format.keep_with_next = True
            location = doc.add_paragraph(f"来源位置：{_source_location(result, segment)}")
            location.paragraph_format.keep_with_next = True
            evidence = segment.get("evidence") or []
            if isinstance(evidence, list):
                doc.add_paragraph("；".join(str(item) for item in evidence))
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_report(result: dict[str, Any], format: str) -> tuple[bytes, str, str]:
    """Return a self-contained UTF-8/OOXML report suitable for downloading."""
    if not isinstance(result, dict):
        raise ValueError("Report result must be an object.")
    if format not in MIME_TYPES:
        raise ValueError("Unsupported document report format.")
    segments = _validated_segments(result)
    if format == "docx":
        data = _docx_report(result, segments)
    elif format == "html":
        data = _html_report(result, segments)
    else:
        payload = dict(result)
        payload["report_export"] = {
            "schema_version": 1,
            "metric_note": REPORT_NOTE,
            "annotation_note": ANNOTATION_NOTE,
            "offset_unit": "Unicode codepoints; start inclusive, end exclusive",
        }
        data = json.dumps(payload, ensure_ascii=False, allow_nan=False, indent=2).encode("utf-8")
    return data, MIME_TYPES[format], _filename(result.get("filename"), format)
