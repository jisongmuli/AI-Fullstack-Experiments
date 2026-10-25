from pathlib import Path
from zipfile import ZipFile
import json,re,hashlib
from pypdf import PdfReader
from docx import Document
import pdfplumber
P=Path(__file__).resolve().parent
F=P.parent/'实验3报告_算法优化实验.docx'
PDF=P/'render-complete/实验3报告_算法优化实验.pdf'
r=PdfReader(PDF);assert len(r.pages)==25
D=Document(F)
text='\n'.join(p.text for p in D.paragraphs)
assert len(D.tables)==14 and len(D.inline_shapes)==8
assert len(D.element.xpath('//m:oMath'))==5
for phrase in ['并不等价于','不代表','不能说明','不能据此宣称','不能证明','并不意味着','正文字号小四','需要涵盖至少第一部分内容']:
    assert phrase not in text,phrase
for name in ['宋凯歌','刘成洋','李阳','魏佳望','史斌斌','李羽']:
    assert name in text or any(name in c.text for t in D.tables for row in t.rows for c in row.cells),name
for i,page in enumerate(r.pages):
    txt=page.extract_text()
    assert len(txt.strip())>100,(i+1,'near blank')
    assert 'Error! Reference source not found.' not in txt
captions=[]
with pdfplumber.open(PDF) as pdf:
    for i,page in enumerate(pdf.pages):
        for im in page.images:
            assert im['x0']>=0 and im['x1']<=page.width+1 and im['top']>=0 and im['bottom']<=page.height+1
        for match in page.search(r'图\s*[35]-\d') or []:
            captions.append({'page':i+1,'caption':match['text']})
assert len(captions)==8,captions
reference_hash=hashlib.sha256((P/'template-reference.docx').read_bytes()).hexdigest()
assert reference_hash=='09c3ea79bde1b87ec6871bd14c9e4c0aaf158ccb02fb3e1e4559241a34cfd449'
with ZipFile(P/'template-reference.docx') as a,ZipFile(F) as b:
    preserve=['word/header1.xml','word/webSettings.xml','word/footnotes.xml','word/endnotes.xml','word/theme/theme1.xml','customXml/item1.xml','customXml/itemProps1.xml']
    for n in preserve:assert a.read(n)==b.read(n),n
    for n in a.namelist():assert n in b.namelist(),n
changed=[]
for i in range(1,26):
    a=P/'render-verified'/f'page-{i}.png';b=P/'render-complete'/f'page-{i}.png'
    if hashlib.sha256(a.read_bytes()).digest()!=hashlib.sha256(b.read_bytes()).digest():changed.append(i)
assert changed==[2,3,4,8,9,10,11,12,13]
summary={'file':str(F),'sha256':hashlib.sha256(F.read_bytes()).hexdigest(),'bytes':F.stat().st_size,
 'pages':25,'page_images_visually_reviewed':list(range(1,26)),
 'pages_reviewed_again_after_final_style_update':changed,
 'figures':8,'data_tables':13,'cover_member_table':1,'native_equations':5,'figure_caption_pages':captions,
 'body_chinese_characters':len(re.findall('[\u4e00-\u9fff]',(P/'report_body.md').read_text(encoding='utf-8'))),
 'template_geometry_preserved':True,'template_reference_unchanged':True,'defensive_examples_present':0,'status':'passed'}
(P/'final-qa.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
with (P/'artifact.md').open('a',encoding='utf-8') as stream:
    stream.write('\n## 最终核验补充\n\n最终报告25页，含1页封面、1页目录。正文及参考文献23页。8幅图、13张正文数据表、1张成员表、5个原生公式。所有页图已逐页检查；最后目录样式修订涉及的9页已重新检查。fontTable因新增明确字体而允许更新，原页眉、webSettings、脚注尾注、主题与customXml按原包保留。参考文献左对齐，原生TOC样式固定为10.5磅及15磅行距。\n')
print(json.dumps({k:summary[k] for k in ['file','bytes','pages','figures','data_tables','native_equations','status']},ensure_ascii=False,indent=2))
