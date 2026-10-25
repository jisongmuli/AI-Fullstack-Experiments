from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
import hashlib,json,re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

P=Path(__file__).resolve().parent
ROOT=P.parents[2]
M=json.loads((P/'measurements.json').read_text(encoding='utf-8'))
REF=P/'template-reference.docx'
FINAL=P.parent/'实验3报告_算法优化实验.docx'
D=Document(REF)
BODY_WIDTH=14.6508

def el(tag,**attrs):
    e=OxmlElement(tag)
    for k,v in attrs.items():e.set(qn(k),str(v))
    return e

def font(run,size=12,bold=None,cn='宋体',en='Times New Roman'):
    run.font.name=en;run.font.size=Pt(size);run.font.color.rgb=RGBColor(0,0,0)
    if bold is not None:run.bold=bold
    rp=run._element.get_or_add_rPr();rf=rp.find(qn('w:rFonts'))
    if rf is None:rf=el('w:rFonts');rp.insert(0,rf)
    for name,value in [('w:eastAsia',cn),('w:ascii',en),('w:hAnsi',en),('w:cs',en)]:rf.set(qn(name),value)
    for key in ['w:asciiTheme','w:eastAsiaTheme','w:hAnsiTheme','w:cstheme']:
        if key in rf.attrib:del rf.attrib[key]
    for tag in ['w:highlight','w:shd','w:color']:
        for node in list(rp.findall(qn(tag))):rp.remove(node)
    rp.append(el('w:color',**{'w:val':'000000'}))
    return run

def style_font(style,size=12,bold=False):
    style.font.name='Times New Roman';style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor(0,0,0)
    rp=style.element.get_or_add_rPr();rf=rp.find(qn('w:rFonts'))
    if rf is None:rf=el('w:rFonts');rp.insert(0,rf)
    rf.set(qn('w:eastAsia'),'宋体');rf.set(qn('w:ascii'),'Times New Roman');rf.set(qn('w:hAnsi'),'Times New Roman')
    for a in list(rf.attrib):
        if a.endswith('Theme'):del rf.attrib[a]

def no_num(p):
    pp=p._p.get_or_add_pPr()
    for x in list(pp.findall(qn('w:numPr'))):pp.remove(x)
    np=el('w:numPr');np.append(el('w:numId',**{'w:val':'0'}));pp.append(np)

def body_format(p):
    pf=p.paragraph_format;pf.first_line_indent=Pt(24);pf.left_indent=Pt(0);pf.right_indent=Pt(0)
    pf.line_spacing=Pt(20);pf.line_spacing_rule=WD_LINE_SPACING.EXACTLY
    pf.space_after=Pt(3);pf.space_before=Pt(0);pf.widow_control=True
    p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    pp=p._p.get_or_add_pPr();pp.append(el('w:snapToGrid',**{'w:val':'0'}))
    no_num(p)

normal=D.styles['Normal'];style_font(normal)
normal.paragraph_format.line_spacing=Pt(20);normal.paragraph_format.space_after=Pt(3)
normal.paragraph_format.first_line_indent=Pt(24)
normal.element.get_or_add_pPr().append(el('w:snapToGrid',**{'w:val':'0'}))
if 'Title' not in D.styles:D.styles.add_style('Title',WD_STYLE_TYPE.PARAGRAPH)
style_font(D.styles['Title'],26,True)
for name,size in [('Heading 1',14),('Heading 2',12),('Heading 3',12)]:
    s=D.styles[name];style_font(s,size,True)
    pf=s.paragraph_format;pf.first_line_indent=Pt(0);pf.line_spacing=Pt(20)
    pf.space_before=Pt(10 if name=='Heading 1' else 7);pf.space_after=Pt(5)
    pf.keep_with_next=True;pf.keep_together=True
    # Source template's built-in heading numbering is replaced by visible labels.
    pp=s.element.get_or_add_pPr()
    for tag in ['w:numPr','w:pBdr']:
        for node in list(pp.findall(qn(tag))):pp.remove(node)
for name,size in [('Exp3 Caption',10.5),('Exp3 Code',9),('Exp3 Reference',10.5),('Exp3 TOC Title',14)]:
    if name not in D.styles:D.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    s=D.styles[name];style_font(s,size,name=='Exp3 TOC Title')
    pf=s.paragraph_format;pf.first_line_indent=Pt(0);pf.line_spacing=Pt(16 if size<12 else 20)
    pf.space_after=Pt(4);pf.space_before=Pt(3)
for level in [1,2]:
    name=f'toc {level}'
    if name not in D.styles:D.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH,builtin=True)
    s=D.styles[name];style_font(s,10.5,level==1)
    s.paragraph_format.first_line_indent=Pt(0);s.paragraph_format.line_spacing=Pt(16)
    s.paragraph_format.space_after=Pt(1);s.paragraph_format.space_before=Pt(0)
    s.paragraph_format.left_indent=Cm(0 if level==1 else .55)
    s.paragraph_format.line_spacing=Pt(15);s.paragraph_format.space_after=Pt(0)
    s.paragraph_format.keep_together=True
    s.element.get_or_add_pPr().append(el('w:snapToGrid',**{'w:val':'0'}))

# Retain the source cover and replace the template's instructional body.
paras=D.paragraphs
date_p=paras[14]
body=D._element.body;cut=False
for child in list(body):
    if cut and child.tag!=qn('w:sectPr'):body.remove(child)
    if child is date_p._p:cut=True
paras=D.paragraphs
paras[4].style=D.styles['Title']
for r in paras[4].runs:font(r,26,True,cn='黑体');r.underline=False
topic=paras[6];topic.clear();topic.style=D.styles['Normal'];no_num(topic)
topic.alignment=WD_ALIGN_PARAGRAPH.CENTER;topic.paragraph_format.first_line_indent=Pt(0)
font(topic.add_run('题目  算法优化实验'),20,True)
subtitle=paras[7];subtitle.clear();subtitle.alignment=WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.first_line_indent=Pt(0);font(subtitle.add_run('多模态AIGC检测模型轻量化与推理验证'),12,False)
for index in [8,9,10,11]:
    paras[index].paragraph_format.first_line_indent=Pt(0)
    paras[index].paragraph_format.line_spacing=Pt(16)
    paras[index].paragraph_format.space_before=Pt(0);paras[index].paragraph_format.space_after=Pt(0)
paras[12].clear();font(paras[12].add_run('组    号    第四组'),14,True);paras[12].paragraph_format.first_line_indent=Pt(0)
paras[13].clear();font(paras[13].add_run('组员信息'),14,True);paras[13].paragraph_format.first_line_indent=Pt(0)
date_p.clear();font(date_p.add_run('时    间    2026年10月25日'),14,True);date_p.paragraph_format.first_line_indent=Pt(0)
members=[
 ('2610321182','宋凯歌','总体设计 结果核验与报告整合'),
 ('2610320011','刘成洋','参数 存储与内存统计'),
 ('2610320065','李阳','文本基线核查与质量指标整理'),
 ('2610321108','魏佳望','特征压缩 长度测量与文本分析'),
 ('2610320004','史斌斌','图像结构 导出与一致性验证'),
 ('2610320063','李羽','批量基准 计算量图表与测试整理')]
cover=D.tables[0]
while len(cover.rows)<7:cover.add_row()
for row,values in zip(cover.rows,[('学号','姓名','分工')]+members):
    row.height=None
    trpr=row._tr.get_or_add_trPr()
    for h in list(trpr.findall(qn('w:trHeight'))):trpr.remove(h)
    for c,text in zip(row.cells,values):
        c.text=text;c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in c.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.line_spacing=Pt(16);p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
            for r in p.runs:font(r,10.5,row is cover.rows[0])
for row in cover.rows:
    for c,w in zip(row.cells,[3.1,2.6,8.74]):c.width=Cm(w)
for c,w in zip(cover.columns,[3.1,2.6,8.74]):c.width=Cm(w)
cover.autofit=False

def ptext(text,style=None):
    p=D.add_paragraph(style=style);body_format(p)
    if style=='Exp3 Reference':
        p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=Pt(16)
        p.paragraph_format.space_after=Pt(7)
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    font(p.add_run(text),10.5 if style=='Exp3 Reference' else 12)
    return p

D.add_page_break()
p=D.add_paragraph('目录',style='Exp3 TOC Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.first_line_indent=Pt(0)
p=D.add_paragraph();p.paragraph_format.first_line_indent=Pt(0)
run=p.add_run();run._r.append(el('w:fldChar',**{'w:fldCharType':'begin'}))
instruction=el('w:instrText');instruction.set(qn('xml:space'),'preserve');instruction.text=' TOC \\o "1-2" \\h \\z \\u '
run._r.append(instruction);run._r.append(el('w:fldChar',**{'w:fldCharType':'separate'}));run._r.append(el('w:t'))
run._r.append(el('w:fldChar',**{'w:fldCharType':'end'}))
D.add_page_break()

def caption(text,table=False):
    p=D.add_paragraph(style='Exp3 Caption');p.paragraph_format.first_line_indent=Pt(0)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=table
    p.paragraph_format.keep_together=True;font(p.add_run(text),10.5)
    return p

def table(headers,rows,widths=None):
    t=D.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths is None:widths=[BODY_WIDTH/len(headers)]*len(headers)
    for cell,text in zip(t.rows[0].cells,headers):cell.text=str(text)
    for values in rows:
        for c,text in zip(t.add_row().cells,values):c.text=str(text)
    props=t._tbl.tblPr
    borders=el('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:borders.append(el('w:'+edge,**{'w:val':'single','w:sz':'4','w:color':'D9D9D9'}))
    props.append(borders)
    margins=el('w:tblCellMar')
    for edge,val in [('top',65),('bottom',65),('left',85),('right',85)]:margins.append(el('w:'+edge,**{'w:w':val,'w:type':'dxa'}))
    props.append(margins)
    t.rows[0]._tr.get_or_add_trPr().append(el('w:tblHeader'))
    for row_id,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(el('w:cantSplit'))
        for j,(cell,width) in enumerate(zip(row.cells,widths)):
            cell.width=Cm(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_id==0:cell._tc.get_or_add_tcPr().append(el('w:shd',**{'w:fill':'E8EDF0'}))
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.line_spacing=Pt(15)
                p.paragraph_format.keep_with_next=row_id==0 or row_id==len(t.rows)-2
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER if row_id==0 or len(cell.text)<16 else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:font(r,10.5,row_id==0)
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    D.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

TABLES={
'environment':(['项目','配置与设置'],[
 ('操作系统','Windows 11 64位'),('处理器','AMD Ryzen 7 9700X  8核16线程'),('物理内存','16262721536字节  约15.15 GiB'),
 ('显卡与计算路径','AMD Radeon RX 9070 GRE  本实验使用CPU'),('Python','3.12.2  项目独立虚拟环境'),('主要框架','torch 2.14.0+cpu  scikit-learn 1.9.1'),
 ('固定配置','随机种子42  算子及数值库线程4'),('复测日期','2026年10月25日')],[3.3,11.35]),
'taskmap':(['任务方向','本实验实现','评价内容'],[
 ('资源占用优化 3.1.5','字符词表30000维压至10000维','文件 参数 数值数组 进程内存'),
 ('推理效率比较 3.1.4','图像原生与TorchScript同机对照','P50 P95 批量吞吐量'),
 ('质量控制','沿用验证集选择规则与原测试记录','宏平均F1 分类错误分布'),
 ('部署验证 3.2','CPU推理与两种执行版本核验','端到端耗时 输出一致性'),
 ('后续引擎路径','ONNX Runtime TensorRT OpenVINO适配分析','算子 线程 硬件与输入约束')],[3.1,5.9,5.65]),
'datasets':(['分支与划分','数量','实验作用'],[
 ('图像训练','4000张','实验二模型训练'),('图像验证','800张','选择最佳训练轮次'),('图像测试','800张','原实验分类质量评价'),
 ('文本原始训练','4855条','字符词表与分类器拟合'),('文本增强训练','599条','训练输入补充'),('文本验证','607条','候选选择与分数校准'),
 ('文本同源测试','607条','原实验最终质量评价'),('补充作文数据','2420条','来源和标签质量分析')],[5.2,2.1,7.35]),
'protocol':(['项目','输入与采样','记录指标'],[
 ('模型主基准','图像32×32 文本固定英文段\nbatch 1与32 各3轮','P50 P95 均值 吞吐量'),
 ('单样本计时','每轮200次 每方法合计600次','完整耗时数组'),('批量计时','每轮80次 每方法合计240次','每批耗时及样本吞吐'),
 ('端到端图像','PNG内存字节 每路径200次','解码至softmax全流程'),('文本长度','100 330 580字符 各150次','长度相关P50与P95'),
 ('连续内存','预热50次后推理1000次 每10次采样','工作集峰值及100个稳定点'),('图像一致性','128个随机输入与2张教学图片','最大logit差及类别差异数')],[3.2,6.3,5.15]),
'storage':(['指标','完整或原生','压缩或导出','变化'],[
 ('文本词表维度','30000','10000','减少66.67%'),('文本分类器参数','30001','10001','减少66.66%'),('文本IDF统计量','30000','10000','减少66.67%'),
 ('文本数值数组 字节','480008','160008','减少66.67%'),('文本文件 字节','479938','167326','减少65.14%'),('固定文本非零特征','418','369','减少11.72%'),
 ('图像可学习参数','140130','140130','相同'),('图像MAC','29196544','29196544','相同'),('图像文件 字节','571026','597526','增加4.64%')],[4.6,3.3,3.3,3.45]),
'quality':(['对象','数据及阶段','指标','原实验记录'],[
 ('完整文本模型','607条验证 未校准','宏平均F1','97.69%'),('压缩文本模型','607条验证 未校准','宏平均F1','97.36%'),
 ('压缩文本模型','607条同源测试 校准后','准确率 宏平均F1','97.69% 97.69%'),('图像CNN','800张测试','准确率 宏平均F1','89.75% 约89.75%'),('图像CNN','800张测试','ROC-AUC','96.30%')],[3.0,4.8,3.15,3.7]),
'memory':(['指标','完整模型 MiB','压缩模型 MiB'],[
 ('基础库基线',f"{M['memory'][0]['baseline_rss']/1024**2:.2f}",f"{M['memory'][1]['baseline_rss']/1024**2:.2f}"),
 ('加载后工作集',f"{M['memory'][0]['after_load_rss']/1024**2:.2f}",f"{M['memory'][1]['after_load_rss']/1024**2:.2f}"),
 ('加载增量',f"{M['memory'][0]['load_delta_rss']/1024**2:.2f}",f"{M['memory'][1]['load_delta_rss']/1024**2:.2f}"),
 ('进程峰值',f"{M['memory'][0]['peak_working_set']/1024**2:.2f}",f"{M['memory'][1]['peak_working_set']/1024**2:.2f}"),
 ('连续调用次数','1000','1000'),('稳定阶段采样点','100','100'),('采样工作集波动','0.00%','0.00%')],[6.25,4.2,4.2]),
'pipeline':(['项目','原生路径','导出路径'],[
 ('内存字节全流程P50 ms','0.57375','0.48560'),('内存字节全流程P95 ms','0.83257','0.69995'),
 ('最大绝对logit差','两版共130个输入','0'),('argmax差异数量','128随机输入 加2张图片','0'),('参数及权重','相同模型状态','导出保存')],[6.25,4.2,4.2]),
'targets':(['评价项目','目标或比较规则','本实验结果'],[
 ('文本文件体积','至少减少40%','减少65.14% 达到'),('验证集宏平均F1','下降不超过1个百分点','下降0.33个百分点 达到'),
 ('文本进程峰值','减少至少30%','减少8.52%'),('连续1000次调用','采样内存波动不超过5%','0.00% 达到'),
 ('图像单张延迟','下降至少40%','P50下降18.87%'),('图像单张吞吐','提升至少50%','平均吞吐提升21.65%'),
 ('导出数值一致性','检查两版输出差异','130输入 最大logit差0')],[4.7,5.0,4.95]),
'evidence':(['资料','主要内容','用途'],[
 ('实验二与实验三说明','质量摘要 原始基准 选择规则','历史实验依据'),('architecture.py','卷积与线性层结构','参数和MAC核算'),
 ('train_and_optimize.py','训练 候选比较 导出与计时','实现流程核查'),('measure_existing_models.py','只读模型复测脚本','本机运行复查'),
 ('measurements.json','环境 指纹 原始耗时 内存轨迹','本机数值来源'),('figures目录','由测量记录生成的8幅图','结果可视化')],[5.4,5.5,3.75])
}
layers=[x for x in M['image_model']['layers'] if x['type'] in ['Conv2d','BatchNorm2d','Linear']]
layer_rows=[];conv_i=bn_i=0
for x in layers:
    if x['type']=='Conv2d':conv_i+=1;name=f'卷积{conv_i}'
    elif x['type']=='BatchNorm2d':bn_i+=1;name=f'批归一化{bn_i}'
    else:name='线性分类'
    layer_rows.append((name,'×'.join(map(str,x['output'][1:])),str(x['parameters']),f"{x['macs']/1e6:.6f}"))
layer_rows.append(('合计','二分类输出','140130','29.196544'))
TABLES['layers']=(['层','输出形状','可学习参数','MAC 百万'],layer_rows,[3.2,4.0,3.55,3.9])
for fam,key,names in [('image','imagebench',{'eager':'原生','torchscript':'TorchScript'}),('text','textbench',{'full':'完整词表','compact':'压缩词表'})]:
    rows=[]
    for x in M['summary']:
        if x['family']==fam:rows.append((names[x['method']],str(x['batch']),f"{x['p50_ms']:.5f}",f"{x['p95_ms']:.5f}",f"{x['throughput_samples_s']:.2f}"))
    TABLES[key]=(['方法','批大小','P50 ms','P95 ms','样本每秒'],rows,[3.25,1.55,3.0,3.0,3.85])

# Native editable Office Math equations.
def mr(text):
    r=el('m:r');pr=el('m:rPr');pr.append(el('m:sty',**{'m:val':'p'}));r.append(pr);t=el('m:t');t.text=text;r.append(t);return r
def sub(base,index):
    s=el('m:sSub');e=el('m:e');e.append(mr(base));i=el('m:sub');i.append(mr(index));s.extend([e,i]);return s
def frac(num,den):
    f=el('m:f');n=el('m:num');d=el('m:den')
    for a in num:n.append(a)
    for a in den:d.append(a)
    f.extend([n,d]);return f
def sup(base,index):
    s=el('m:sSup');e=el('m:e');e.append(mr(base));i=el('m:sup');i.append(mr(index));s.extend([e,i]);return s
def equation(key):
    p=D.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
    p.paragraph_format.line_spacing=1
    p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(7);p.paragraph_format.keep_together=True
    om=el('m:oMathPara');math=el('m:oMath')
    if key=='conv':nodes=[mr('MAC = '),sub('H','o'),mr(' × '),sub('W','o'),mr(' × '),sub('C','o'),mr(' × '),frac([sub('C','i')],[mr('g')]),mr(' × '),sub('k','h'),mr(' × '),sub('k','w')]
    elif key=='idf':nodes=[mr('idf(t) = log('),frac([mr('1 + N')],[mr('1 + df(t)')]),mr(') + 1')]
    elif key=='linear':nodes=[mr('z = '),sup('w','T'),mr('x + b     P = d + 1')]
    elif key=='reduction':nodes=[mr('R = '),frac([sub('L','before'),mr(' − '),sub('L','after')],[sub('L','before')]),mr(' × 100%')]
    elif key=='throughput':nodes=[mr('Q = '),frac([mr('B × 1000')],[sub('L','mean')]),mr('   单位 样本每秒')]
    for n in nodes:math.append(n)
    om.append(math);p._p.append(om)

in_code=False;code_lines=[]
for line in (P/'report_body.md').read_text(encoding='utf-8').splitlines():
    if line.startswith('```'):
        if not in_code:in_code=True;code_lines=[]
        else:
            for i,text in enumerate(code_lines):
                p=D.add_paragraph(style='Exp3 Code');p.paragraph_format.first_line_indent=Pt(0)
                p.paragraph_format.line_spacing=Pt(13);p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.keep_with_next=i<len(code_lines)-1
                font(p.add_run(text),9,False,cn='宋体',en='Consolas')
            D.add_paragraph().paragraph_format.space_after=Pt(0)
            in_code=False
        continue
    if in_code:code_lines.append(line);continue
    if not line.strip():continue
    if line.startswith('# '):
        text=line[2:];p=D.add_paragraph(text,style='Heading 1');no_num(p)
        p.paragraph_format.first_line_indent=Pt(0)
        p.paragraph_format.page_break_before=False
        for r in p.runs:font(r,14,True)
    elif line.startswith('## '):
        p=D.add_paragraph(line[3:],style='Heading 2');no_num(p);p.paragraph_format.first_line_indent=Pt(0)
        for r in p.runs:font(r,12,True)
    elif line.startswith('@TABLE '):
        key,title=line[7:].split('|',1);caption(title,True);table(*TABLES[key])
    elif line.startswith('@FIG '):
        filename,title=line[5:].split('|',1)
        p=D.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
        p.paragraph_format.line_spacing=1
        p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(2)
        r=p.add_run();r.add_picture(str(P/'figures'/filename),width=Cm(14.3 if filename!='confusion.png' else 11.8))
        inline=r._r.xpath('.//wp:docPr')[0];inline.set('descr',title)
        caption(title)
    elif line.startswith('@EQ '):equation(line[4:])
    else:ptext(line,'Exp3 Reference' if re.match(r'^\[\d+\]',line) else None)

sec=D.sections[0]
sec.different_first_page_header_footer=True
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.paragraph_format.first_line_indent=Pt(0)
footer.paragraph_format.line_spacing=Pt(14)
font(footer.add_run('第 '),10.5)
for field_name,suffix in [('PAGE',' 页  共 '),('NUMPAGES',' 页')]:
    fld=el('w:fldSimple',**{'w:instr':field_name});r=el('w:r');t=el('w:t');t.text='1';r.append(t);fld.append(r);footer._p.append(fld)
    font(footer.add_run(suffix),10.5)
sec.first_page_footer.paragraphs[0].clear()
settings=D.settings.element
existing=settings.find(qn('w:updateFields'))
if existing is None:settings.append(el('w:updateFields',**{'w:val':'true'}))
else:existing.set(qn('w:val'),'true')
D.core_properties.title='算法优化实验报告'
D.core_properties.subject='多模态AIGC检测模型轻量化与推理验证'
D.core_properties.author='第四组'
D.core_properties.keywords='算法优化 特征裁剪 模型压缩 推理基准 资源占用'
D.save(FINAL)
print(json.dumps({'output':str(FINAL),'paragraphs':len(D.paragraphs),'tables':len(D.tables),'figures':len(D.inline_shapes),'native_equations':len(D.element.xpath('//m:oMath')),'reference_sha256':hashlib.sha256(REF.read_bytes()).hexdigest()},ensure_ascii=False))
