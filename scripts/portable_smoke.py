"""Runtime smoke test using generated fixtures; no accuracy measurement."""
from __future__ import annotations
import argparse,io,json,time,sys
from pathlib import Path
import httpx
from PIL import Image
from docx import Document
import pymupdf

p=argparse.ArgumentParser()
p.add_argument('--url',default='http://127.0.0.1:8003')
p.add_argument('--output',type=Path)
a=p.parse_args()
root=Path(__file__).resolve().parents[1]
record={'purpose':'portable package workflow smoke test, not detection accuracy',
        'accuracy_evaluated':False,'checks':[],'python':sys.version.split()[0]}
ZH='在一次公开数据整理过程中，小组先查看文件的格式和来源，再检查样本是否存在缺失与重复。处理完成后，每个文件都保存稳定的样本编号，并将处理参数记录在清单中。这样既方便其他同学复现步骤，也能够在发现问题时回到原始记录。'
EN='The team checked the source files before preparing the data. Each record received a stable identifier. Missing values and duplicate records were counted separately, and only the training split was used to fit numeric transformations.'
doc=Document();doc.add_paragraph(ZH);doc.add_paragraph(EN);doc.add_paragraph('Emoji offset test 😀 with ordinary punctuation.')
docbytes=io.BytesIO();doc.save(docbytes)
pdf=pymupdf.open();page=pdf.new_page();page.insert_textbox((50,50,500,750),EN,fontsize=12);pdfbytes=pdf.tobytes();pdf.close()
im=Image.new('RGB',(256,256))
im.putdata([(x,y,(x+y)//2) for y in range(256) for x in range(256)])
imbytes=io.BytesIO();im.save(imbytes,format='PNG')
with httpx.Client(base_url=a.url,timeout=180) as client:
    def check(name,method,path,**kw):
        started=time.perf_counter();response=client.request(method,path,**kw)
        response.raise_for_status()
        record['checks'].append({'name':name,'http_status':response.status_code,
                                 'latency_ms':round((time.perf_counter()-started)*1000,1)})
        return response
    assert check('health','GET','/health').json()['status']=='ok'
    home=check('homepage','GET','/')
    assert 'data-view="document"' in home.text
    assert 'no-store' in home.headers['cache-control']
    for language,text in [('Chinese',ZH),('English',EN)]:
        result=check(language+' text','POST','/api/text/detect',json={'text':text}).json()
        assert result['segments'] and any(s.get('risk_score') is not None for s in result['segments'])
    original=check('NTIRE 2026 source image','POST','/api/image/detect',files={'file':('sample-2026.jpg',(root/'web/assets/sample-real-2026.jpg').read_bytes(),'image/jpeg')}).json()
    assert original['risk_score'] is not None
    ordinary=check('ordinary image dual models','POST','/api/image/detect',files={'file':('generated-gradient-fixture.png',imbytes.getvalue(),'image/png')}).json()
    assert len(ordinary['model_scores'])>=1
    record['checks'][-1]['scores_are_not_ground_truth']=True
    ids=[]
    for name,data,mime in [('fixture.docx',docbytes.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document'),('fixture.pdf',pdfbytes,'application/pdf'),('fixture.txt',(ZH+'\n\n'+EN).encode(),'text/plain')]:
        r=check(name+' detection','POST','/api/document/detect',files={'file':(name,data,mime)}).json()
        assert r['detection']['segments'] and r['extracted_text']
        ids.append(r['report_id'])
    for fmt in ['docx','html','json']:
        response=check(fmt+' export','GET',f'/api/document/reports/{ids[0]}/{fmt}')
        if fmt=='docx':assert response.content[:2]==b'PK'
        elif fmt=='html':assert '<html' in response.text.lower()
        else:assert response.json()['report_id']==ids[0]
    available=check('available report list','GET','/api/document/reports').json()['reports']
    assert set(ids).issubset({item['report_id'] for item in available})
    assert all('extracted_text' not in item and 'detection' not in item for item in available)
    preview=check('local LoRA editing preview','POST','/api/text/edit/llm-preview',json={'text':'She go to school yesterday and learn a new lesson.'}).json()
    assert preview['status'] in {'rejected','review_candidate'}
    record['checks'][-1]['screened_status']=preview['status']
    for rid in ids:check('delete temporary report','DELETE',f'/api/document/reports/{rid}')
record['status']='passed'
if a.output:
    a.output.parent.mkdir(parents=True,exist_ok=True)
    a.output.write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(record,ensure_ascii=False,indent=2))
